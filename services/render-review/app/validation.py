"""Compiler validation gates (PDF section 5.3).

Re-open the compiled PPTX bytes with python-pptx and verify:
  - PPTX opens successfully.
  - Mandatory slide title and legend exist.
  - Shape counts are consistent with the layout model.
"""
from __future__ import annotations

import io

from pptx import Presentation


class ValidationFailure(Exception):
    def __init__(self, errors: list[str]):
        self.errors = errors
        super().__init__("; ".join(errors))


def validate_pptx(pptx_bytes: bytes, layout: dict) -> dict:
    """Return stats dict; raise ValidationFailure on any gate failure."""
    errors: list[str] = []

    # Gate 1: opens successfully.
    try:
        prs = Presentation(io.BytesIO(pptx_bytes))
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailure([f"PPTX failed to open: {exc}"]) from exc

    if len(prs.slides) < 1:
        raise ValidationFailure(["PPTX has no slides"])
    slide = prs.slides[0]

    names = [sh.name for sh in slide.shapes]

    # Gate 2: title + legend present.
    if not any(n == "slide-title" for n in names):
        errors.append("mandatory slide title is missing")
    if not any(n == "legend-title" for n in names):
        errors.append("mandatory legend is missing")

    # Gate 3: shape counts match the layout.
    n_objects = len(layout.get("objects", []))
    container_shapes = sum(1 for n in names if n.startswith("container:"))
    icon_shapes = sum(1 for n in names if n.startswith("icon:"))
    drawn_objects = container_shapes + icon_shapes
    if drawn_objects != n_objects:
        errors.append(
            f"object shape count {drawn_objects} != layout objects {n_objects}"
        )

    # Every layout object id should be present as a shape.
    drawn_ids = {
        n.split(":", 1)[1]
        for n in names
        if n.startswith(("container:", "icon:"))
    }
    for o in layout.get("objects", []):
        if o["objectId"] not in drawn_ids:
            errors.append(f"object {o['objectId']} not drawn")

    if errors:
        raise ValidationFailure(errors)

    return {
        "slides": len(prs.slides),
        "containerShapes": container_shapes,
        "iconShapes": icon_shapes,
        "connectorSegments": sum(1 for n in names if n.startswith("conn:")),
        "totalShapes": len(names),
    }
